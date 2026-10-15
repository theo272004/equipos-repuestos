#!/usr/bin/env python3
"""Pone al día el plan de mantenimiento de la app con el Excel "MANTENIMIENTO POR SISTEMAS".

Lo que la app lee del Excel está en assets/equipos.js (window.EQUIPOS_PLAN): por cada
pieza de cada equipo, la frecuencia, la última ejecución, la próxima y el estado
(xls.f, xls.ej, xls.px, xls.st). Este script lee la hoja "PLAN MTTO " y:

  1. Actualiza esas cuatro cifras en las filas que ya existen (se cruzan por código
     del equipo + código del repuesto + sistema).
  2. Agrega las filas que el Excel trae y la app no tenía (cantidad, existencias por
     bodega y observaciones salen del Excel; el precio y la ubicación se copian de
     otra fila de la app con el mismo código).
  3. No borra nada: si el Excel dejó de listar una pieza, la app la conserva (con sus
     notas y precio) y solo deja de actualizarla.
  4. Sube la fecha de corte ("FECHA ACTUAL" del Excel).

Uso:
  python3 scripts/actualizar-plan-mtto.py "ruta/MANTENIMIENTO POR SISTEMAS.xlsm"
  python3 scripts/actualizar-plan-mtto.py "ruta/archivo.xlsm" --probar     (no escribe, solo cuenta)

Es seguro correrlo varias veces: con el mismo Excel la segunda vez no cambia nada.
No ejecuta las macros del archivo: solo lee los valores guardados.
"""
import json
import re
import sys
import warnings
from collections import defaultdict
from pathlib import Path

import openpyxl

warnings.filterwarnings("ignore")
RAIZ = Path(__file__).resolve().parent.parent
DESTINO = RAIZ / "assets" / "equipos.js"


def iso(v):
    return v.date().isoformat() if hasattr(v, "date") else (str(v).strip() if v not in (None, "") else "")


def cod(v):
    if v in (None, ""):
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def fecha(v):
    t = iso(v)
    return t if re.fullmatch(r"\d{4}-\d{2}-\d{2}", t) else ""


def num(v):
    return int(v) if isinstance(v, (int, float)) and not isinstance(v, bool) and float(v).is_integer() else (v if isinstance(v, (int, float)) else 0)


def texto(v):
    t = "" if v is None else str(v).strip()
    return "" if t.startswith("#") else t


def norm(t):
    return re.sub(r"\s+", " ", str(t or "").upper()).strip()


def leer_excel(ruta):
    wb = openpyxl.load_workbook(ruta, data_only=True)
    hoja = next(ws for ws in wb.worksheets if ws.title.strip().upper() == "PLAN MTTO")
    corte = fecha(hoja["X1"].value)
    filas = []
    for r in hoja.iter_rows(min_row=3, max_row=hoja.max_row, values_only=True):
        eq = cod(r[0])
        if not eq:
            continue
        filas.append({
            "eq": eq, "sis": texto(r[3]), "act": texto(r[4]), "cod": cod(r[5]), "d": texto(r[6]),
            "q": num(r[9]), "e": num(r[10]), "w": [num(r[11]), num(r[12]), num(r[13]), num(r[14])],
            "f": texto(r[17]), "ej": fecha(r[18]), "px": fecha(r[19]), "st": texto(r[21]), "o": texto(r[22]),
        })
    return filas, corte


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if not args:
        sys.exit(__doc__)
    probar = "--probar" in sys.argv
    filas, corte = leer_excel(args[0])

    s = DESTINO.read_text(encoding="utf-8")
    m = re.search(r"window\.EQUIPOS_PLAN = ", s)
    cuerpo = s[m.end():]
    fin = cuerpo.rstrip().rfind(";")
    datos = json.loads(cuerpo[:fin])
    por_eq = {str(e["c"]): e for e in datos["equipos"]}

    # Para copiar precio y ubicación de una pieza que la app ya conoce
    conocidas = {}
    for e in datos["equipos"]:
        for r in e["r"]:
            if r.get("cod"):
                conocidas.setdefault(r["cod"], r)

    usados = set()
    cambiadas, agregadas, sin_equipo = [], [], 0
    for x in filas:
        e = por_eq.get(x["eq"])
        if not e:
            sin_equipo += 1
            continue
        # Misma pieza: mismo código y mismo sistema; si no, el mismo código
        cand = [i for i, r in enumerate(e["r"]) if (x["eq"], i) not in usados and r.get("cod", "") == x["cod"]]
        mismo = [i for i in cand if norm(e["r"][i].get("s")) == norm(x["sis"])]
        pool = mismo or cand
        nuevo_xls = {"f": x["f"], "ej": x["ej"], "px": x["px"], "st": x["st"]}
        # Con filas idénticas (mismo código y sistema) se elige la que ya dice lo mismo: no se voltean entre sí
        i = next((k for k in pool if e["r"][k].get("xls") == nuevo_xls), (pool or [None])[0])
        if i is None:
            if not x["cod"]:
                continue
            base = conocidas.get(x["cod"], {})
            e["r"].append({
                "s": x["sis"], "a": x["act"], "cod": x["cod"], "d": x["d"] or base.get("d", ""),
                "q": x["q"], "e": x["e"], "w": x["w"], "o": x["o"] or base.get("o", ""),
                "ub": base.get("ub", ""), "pu": base.get("pu", 0),
                "xls": {"f": x["f"], "ej": x["ej"], "px": x["px"], "st": x["st"]},
            })
            usados.add((x["eq"], len(e["r"]) - 1))
            agregadas.append((e["n"], x["cod"], x["sis"], x["ej"]))
            continue
        usados.add((x["eq"], i))
        r = e["r"][i]
        nuevo = {"f": x["f"], "ej": x["ej"], "px": x["px"], "st": x["st"]}
        if r.get("xls") != nuevo:
            cambiadas.append((e["n"], r.get("cod"), (r.get("d") or "")[:34], r.get("xls", {}).get("ej", ""), x["ej"], x["st"]))
            r["xls"] = nuevo

    print(f"Excel: {len(filas)} filas, corte {corte or '?'}")
    print(f"Piezas con fechas actualizadas: {len(cambiadas)} · agregadas: {len(agregadas)} · de equipos que la app no tiene: {sin_equipo}")
    por_eq_cambio = defaultdict(int)
    for n, *_ in cambiadas:
        por_eq_cambio[n] += 1
    for n, k in sorted(por_eq_cambio.items(), key=lambda t: -t[1]):
        print(f"  {k:3d}  {n}")
    for n, c, sis, ej in agregadas:
        print(f"  + {n} · {c} · {sis or '—'} · ejecutado {ej or '—'}")
    if probar:
        print("(--probar: no se escribió nada)")
        return
    if corte:
        datos["corte"] = corte
    nuevo = json.dumps(datos, ensure_ascii=False, indent=2)
    DESTINO.write_text(s[:m.end()] + nuevo + cuerpo[fin:], encoding="utf-8")
    print(f"Escrito: {DESTINO.relative_to(RAIZ)}")


if __name__ == "__main__":
    main()
