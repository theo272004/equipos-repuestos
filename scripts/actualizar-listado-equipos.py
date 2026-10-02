# Pasa al registro de equipos (assets/equipos.js) los listados oficiales en Excel:
#   DMM-179B  Sede 4    (hoja "Hoja3")
#   DMM-179   Planta 2  (hoja "Hoja3")
#
#   python3 scripts/actualizar-listado-equipos.py DMM-179B_Sede4.xlsx DMM-179_Planta2.xlsx
#
# Que hace con cada fila del listado vigente (Hoja3):
#   - Si el equipo ya esta en el registro (mismo codigo y misma sede), le pone
#     el area ("REFERIDO A" -> ref) y el tipo de area ("TIPO UBICACION" -> zona).
#     No le toca el nombre, ni el centro de costo, ni los repuestos: el nombre ya
#     esta con tildes y bien escrito, y el cc de los equipos con plan viene de
#     MANTENIMIENTO POR SISTEMAS.
#   - Si el codigo no esta, lo agrega como equipo nuevo (sin repuestos), justo
#     despues del equipo que lo precede en el listado.
#
# Las otras hojas del mismo libro son de 2018 (listado y programa anual de
# preventivo). Varios codigos se reutilizaron desde entonces para otro equipo,
# asi que solo se toma lo de 2018 (familia, clase y cuantos preventivos al ano
# se programaron) cuando la descripcion de 2018 es claramente el mismo equipo.
#
# Los Excel no se suben al repositorio: es publico y el pie de firmas trae
# nombres de personas.

import json
import re
import sys
import unicodedata
from difflib import SequenceMatcher
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
RUTA_PLAN = RAIZ / "assets" / "equipos.js"

# Filas del listado vigente que no se aplican tal cual (errores del Excel,
# confirmados contra el listado de 2018 y el PDF firmado).
NO_RENOMBRAR = {
    # En la Hoja3 de Planta 2 los tanques de jarabe quedaron con el nombre del
    # tanque de arriba (arrastre de celda); en 2018 y en el PDF son de jarabe.
    "17303006", "17303007", "17303011", "17303012", "17303019",
    # Banda BLISTEADO#3 aparece como #2, que es otro equipo (17371002).
    "17371003",
}

# Codigos que cambiaron en el listado: el equipo es el mismo, solo se le
# corrige el codigo (no tenian repuestos ni historial colgado del codigo viejo).
CODIGO_NUEVO = {
    "135408024": "135351002",  # CAB. Flujo Laminar Dispensacion 1
    "135408022": "135351003",  # CAB. Flujo Laminar Dispensacion 2
}

# Equipos que ya tienen ficha propia en machines-data.js: el registro los
# enlaza por id para que la ficha rica reciba sus repuestos y su codigo.
FICHA_PROPIA = {
    "17333008": "njp3500",
}

# Nombres bonitos para los equipos nuevos (el Excel los trae en mayusculas).
NOMBRES = {
    "17332011": "Blisteadora #8 UPS 300",
    "17332012": "Blisteadora #9 UPS 300",
    "125334008": "Chequeadora de Peso Insight",
    "125305002": "Codificadora EasyPrint #2",
    "125305003": "Codificadora EasyPrint #3",
    "125305004": "Codificadora EasyPrint #4",
    "17365013": "Desempolvador Vertical NJP3500",
    "17306009": "Detector de Metal #9 NJP 3500",
    "17333008": "Encapsuladora NJP 3500",
}

SEDES = {"SEDE 4": "SEDE 4", "PLANTA 2": "PLANTA 2"}


def sin_tildes(s):
    return unicodedata.normalize("NFD", s).encode("ascii", "ignore").decode()


def limpio(s):
    return re.sub(r"\s+", " ", str(s or "")).strip()


def codigo(v):
    s = limpio(v)
    if re.fullmatch(r"\d+(\.0)?", s):
        return str(int(float(s)))  # el registro guarda los codigos sin ceros a la izquierda
    return s


def zona(s):
    s = sin_tildes(limpio(s)).upper().replace("EREA", "AREA")
    m = re.fullmatch(r"AREA (BLANCA|GRIS|NEGRA)", s)
    return f"ÁREA {m.group(1)}" if m else limpio(s).upper()


def ref(s):
    s = limpio(s).upper()
    for a, b in (("LIQUIDOS", "LÍQUIDOS"), ("SOLIDOS", "SÓLIDOS"), ("PREPARACION", "PREPARACIÓN"),
                 ("AREA TECNICA", "ÁREA TECNICA")):
        s = re.sub(rf"\b{a}\b", b, s)
    return s


def titulo(s):
    chicas = {"de", "del", "la", "las", "el", "los", "para", "en", "y", "con", "a"}
    out = []
    for p in limpio(s).split(" "):
        if re.search(r"[\d#/.]", p) or len(p) <= 3:
            out.append(p if p.lower() not in chicas else p.lower())
        elif p.lower() in chicas:
            out.append(p.lower())
        else:
            out.append(p[:1] + p[1:].lower())
    return " ".join(out)


def leer_vigente(ruta):
    ws = openpyxl.load_workbook(ruta, data_only=True)["Hoja3"]
    filas = []
    for r in ws.iter_rows(min_row=5, values_only=True):
        if not r[1] or not re.fullmatch(r"\d+(\.0)?", limpio(r[1])):
            continue  # cabeceras y pie de firmas
        filas.append({
            "c": codigo(r[1]),
            "tipo": "AUXILIAR" if limpio(r[2]).upper().startswith("AUX") else limpio(r[2]).upper(),
            "n": limpio(r[3]).upper(),
            "u": SEDES.get(limpio(r[4]).upper(), limpio(r[4]).upper()),
            "ref": ref(r[5]),
            "zona": zona(r[6]),
        })
    return filas


def norm_2018(s):
    s = sin_tildes(limpio(s)).upper()
    s = s.replace("TABLETEADOLETEADORA", "TABLETEADORA").replace("LIQUIDOSUIDOS", "LIQUIDOS").replace("LIQUIDOSUIDO", "LIQUIDO")
    s = re.sub(r"^U\. ?MANEJADORA", "UMA", s)
    s = re.sub(r"^U\. ?CONDENSADORA", "UCA", s)
    s = re.sub(r"\(NO ENCONTRAD[OA]\)", "", s)
    s = re.sub(r"\bEN RECUBRIMIENTO.*$", "", s)
    s = re.sub(r"\bDE NUCLEOS INERTES\b", "", s)
    return re.sub(r"[^A-Z0-9]+", "", s)


def leer_2018(ruta):
    """codigo -> {desc, fam, cl, est, pm}. Familia y clase del LISTADO 2018;
    preventivos programados del PM EQUIPOS 2018."""
    wb = openpyxl.load_workbook(ruta, data_only=True)
    datos = {}
    for r in wb["LISTADO DE EQUIPOS 2018"].iter_rows(min_row=9, values_only=True):
        c = codigo(r[3])
        if not c.isdigit():
            continue
        datos.setdefault(c, []).append({
            "desc": limpio(r[7]), "fam": limpio(r[4]), "cl": limpio(r[5]), "est": limpio(r[9]).upper(),
        })
    pm = {}
    for r in wb["PM EQUIPOS 2018"].iter_rows(min_row=9, values_only=True):
        c = codigo(r[3])
        if not c.isdigit():
            continue
        prog = [r[12 + 2 * i] for i in range(12)]
        pm.setdefault(c, []).append((limpio(r[7]), sum(1 for v in prog if isinstance(v, (int, float)) and v > 0)))
    for c, lista in datos.items():
        for d in lista:
            for desc, n in pm.get(c, []):
                if norm_2018(desc) == norm_2018(d["desc"]):
                    d["pm"] = n
    return datos


def de_2018(datos, eq):
    """Lo de 2018 que le corresponde al equipo, si la descripcion coincide."""
    mejor, nota = None, 0
    for d in datos.get(eq["c"], []):
        s = SequenceMatcher(None, norm_2018(d["desc"]), norm_2018(eq["nOriginal"])).ratio()
        if s > nota:
            mejor, nota = d, s
    if not mejor or nota < 0.85:
        return None
    h = {"fam": mejor["fam"], "cl": mejor["cl"]}
    if mejor.get("pm"):
        h["pm"] = mejor["pm"]
    if mejor["est"] and mejor["est"] != "ACTIVO":
        h["est"] = mejor["est"]
    return h


def main(rutas):
    texto = RUTA_PLAN.read_text(encoding="utf8")
    ini, fin = texto.index("{"), texto.rindex("}")
    plan = json.loads(texto[ini:fin + 1])
    equipos = plan["equipos"]

    for eq in equipos:
        if eq["c"] in CODIGO_NUEVO:
            nuevo = CODIGO_NUEVO[eq["c"]]
            print(f"codigo  {eq['c']} -> {nuevo}  {eq['n']}")
            eq["c"], eq["id"] = nuevo, eq["id"].replace(eq["c"], nuevo)
        if eq["u"] == "SE":
            eq["u"] = "SEDE 4"

    def buscar(c, u):
        return next((e for e in equipos if e["c"] == c and e["u"] == u), None)

    datos18 = {}
    nuevos, tocados = [], 0
    for ruta in rutas:
        for c, lista in leer_2018(ruta).items():
            datos18.setdefault(c, lista)
        anterior = None
        for f in leer_vigente(ruta):
            eq = buscar(f["c"], f["u"])
            if eq is None:
                otra_sede = next((e for e in equipos if e["c"] == f["c"]), None)
                mid = FICHA_PROPIA.get(f["c"], "")
                eq = {
                    "c": f["c"],
                    "id": mid or ("eq-" + f["c"] + ("-p2" if otra_sede and f["u"] == "PLANTA 2" else "")),
                    "mid": mid,
                    "n": NOMBRES.get(f["c"]) or titulo(f["n"]),
                    "nOriginal": f["n"],
                    "u": f["u"],
                    "cc": f"{f['ref']} · {f['zona']}",
                    "tipo": f["tipo"],
                    "r": [],
                }
                pos = equipos.index(anterior) + 1 if anterior else len(equipos)
                equipos.insert(pos, eq)
                nuevos.append(eq)
            elif f["c"] not in NO_RENOMBRAR and sin_tildes(eq["nOriginal"]).upper().replace(" ", "") != sin_tildes(f["n"]).replace(" ", ""):
                print(f"nombre  {f['c']} registro={eq['nOriginal']!r} listado={f['n']!r} (se deja el del registro)")
            eq["ref"], eq["zona"] = f["ref"], f["zona"]
            tocados += 1
            anterior = eq

    con18 = 0
    for eq in equipos:
        h = de_2018(datos18, eq)
        if h:
            eq["h18"] = h
            con18 += 1
        else:
            eq.pop("h18", None)

    fuera = [e for e in equipos if "ref" not in e]
    plan["fuente"] = ("Listados oficiales de equipos FARMACAPSULAS (Sede 4: DMM-179B V02-2026, "
                      "Planta 2: DMM-179 V05-2025) + MANTENIMIENTO POR SISTEMAS + listado y programa PM 2018")
    salida = texto[:ini] + json.dumps(plan, indent=2, ensure_ascii=False) + texto[fin + 1:]
    salida = salida.replace("Generado a partir de los PDF oficiales", "Generado a partir de los listados oficiales")
    RUTA_PLAN.write_text(salida, encoding="utf8")

    for e in nuevos:
        print(f"nuevo   {e['c']:>10}  {e['u']:<8}  {e['n']}  ({e['cc']})" + (f"  -> ficha {e['mid']}" if e["mid"] else ""))
    print(f"\n{len(equipos)} equipos en el registro: {tocados} del listado vigente ({len(nuevos)} nuevos), "
          f"{len(fuera)} que no estan en el listado, {con18} con datos de 2018.")
    for e in fuera:
        print(f"  no esta en el listado: {e['c']:>10}  {e['u']:<8}  {e['n']}  ({len(e['r'])} repuestos)")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__ or "uso: actualizar-listado-equipos.py LISTADO.xlsx [LISTADO2.xlsx ...]")
    main(sys.argv[1:])
